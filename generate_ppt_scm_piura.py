from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTPUT_FILE = Path("Presentacion_SCM_Piura_Gerencia_General.pptx")

COLOR_NAVY = RGBColor(15, 55, 110)
COLOR_BLUE = RGBColor(31, 111, 235)
COLOR_SOFT = RGBColor(239, 244, 248)
COLOR_BORDER = RGBColor(210, 223, 235)
COLOR_TEXT = RGBColor(36, 52, 71)
COLOR_MUTED = RGBColor(98, 117, 137)
COLOR_GREEN = RGBColor(25, 135, 84)
COLOR_ORANGE = RGBColor(217, 136, 43)
COLOR_RED = RGBColor(220, 53, 69)
COLOR_WHITE = RGBColor(255, 255, 255)


def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_WHITE


def add_textbox(slide, x, y, w, h, text="", size=18, bold=False,
                color=COLOR_TEXT, align=PP_ALIGN.LEFT, font_name="Calibri"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font_name
    run.font.color.rgb = color
    p.alignment = align
    tf.word_wrap = True
    return box


def add_bullets(slide, x, y, w, h, bullets, size=20, level0_color=COLOR_TEXT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(8)
        p.font.size = Pt(size)
        p.font.name = "Calibri"
        p.font.color.rgb = level0_color
    return box


def add_title(slide, title, subtitle=None):
    add_textbox(slide, Inches(0.6), Inches(0.35), Inches(8.8), Inches(0.55),
                title, size=28, bold=True, color=COLOR_NAVY)
    if subtitle:
        add_textbox(slide, Inches(0.6), Inches(0.88), Inches(8.9), Inches(0.4),
                    subtitle, size=12, color=COLOR_MUTED)


def add_footer(slide, index):
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(7.05), Inches(12.25), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_BORDER
    line.line.fill.background()
    add_textbox(slide, Inches(0.6), Inches(7.08), Inches(8.5), Inches(0.2),
                "Sistema SCM Piura | Presentacion tecnico-gerencial", size=9, color=COLOR_MUTED)
    add_textbox(slide, Inches(12.25), Inches(7.03), Inches(0.45), Inches(0.25),
                str(index), size=10, bold=True, color=COLOR_NAVY, align=PP_ALIGN.RIGHT)


def add_key_message(slide, text):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(6.18), Inches(5.95), Inches(0.55)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_SOFT
    shape.line.color.rgb = COLOR_BORDER
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r1 = p.add_run()
    r1.text = "Mensaje clave: "
    r1.font.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = COLOR_NAVY
    r2 = p.add_run()
    r2.text = text
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_TEXT


def add_visual_placeholder(slide, x, y, w, h, title, items):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_SOFT
    shape.line.color.rgb = COLOR_BLUE
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Pt(10)
    tf.margin_right = Pt(10)
    tf.margin_top = Pt(8)
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.LEFT
    r0 = p0.add_run()
    r0.text = title
    r0.font.bold = True
    r0.font.size = Pt(15)
    r0.font.color.rgb = COLOR_NAVY

    for item in items:
        p = tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(5)

    return shape


def add_kpi_card(slide, x, y, w, h, title, value, color):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_WHITE
    shape.line.color.rgb = COLOR_BORDER
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Pt(12)
    tf.margin_top = Pt(10)
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = title
    r1.font.size = Pt(11)
    r1.font.color.rgb = COLOR_MUTED
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = value
    r2.font.size = Pt(24)
    r2.font.bold = True
    r2.font.color.rgb = color


def add_process_flow(slide):
    steps = [
        "Login y acceso",
        "Carga de padron",
        "Panel de control",
        "Procesar lecturas",
        "Carga manual",
        "Exportacion y cierre",
    ]
    x_positions = [0.7, 2.75, 4.8, 6.85, 8.9, 10.95]
    for idx, step in enumerate(steps):
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(x_positions[idx]), Inches(2.0), Inches(1.55), Inches(0.82)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_SOFT if idx % 2 == 0 else COLOR_WHITE
        shape.line.color.rgb = COLOR_BLUE
        shape.text_frame.text = step
        for p in shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = COLOR_NAVY

        if idx < len(steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(x_positions[idx] + 1.58), Inches(2.26), Inches(0.3), Inches(0.28)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = COLOR_BLUE
            arrow.line.color.rgb = COLOR_BLUE


def add_table_matrix(slide, x, y, w, h, data):
    rows = len(data)
    cols = len(data[0])
    table = slide.shapes.add_table(rows, cols, x, y, w, h).table
    table.first_row = True
    col_widths = [1.55, 2.15, 2.65, 1.9]
    for idx, width in enumerate(col_widths[:cols]):
        table.columns[idx].width = Inches(width)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = data[r][c]
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_SOFT if r == 0 else COLOR_WHITE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if r == 0 else PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.size = Pt(11 if r == 0 else 10.5)
                    run.font.bold = r == 0
                    run.font.color.rgb = COLOR_NAVY if r == 0 else COLOR_TEXT


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slides_data = [
        {
            "title": "Sistema SCM Piura",
            "subtitle": "Plataforma de control operativo para padron, lecturas, trazabilidad y regularizacion",
            "bullets": [
                "Presentacion tecnico-gerencial para Gerencia General.",
                "Desarrollo orientado al control integral del proceso operativo mensual.",
                "Corte de ejemplo mostrado en la interfaz: periodo 2026-04.",
            ],
            "visual": ("Evidencia visual sugerida", [
                "Captura 01: Ingreso al sistema",
                "Uso: portada o slide inicial",
            ]),
            "key": "La solucion integra padrón, lectura, regularizacion y control en una sola plataforma."
        },
        {
            "title": "Situacion inicial y oportunidad de mejora",
            "subtitle": "Problemas operativos que el desarrollo viene resolviendo",
            "bullets": [
                "Procesos distribuidos entre archivos, formatos y revisiones manuales.",
                "Baja visibilidad del avance real por periodo, suministro y cuadrilla.",
                "Mayor riesgo de errores por duplicidad, digitacion y cruce de fuentes.",
                "Tratamiento poco estandarizado de excepciones y lecturas incompletas.",
            ],
            "visual": ("Impacto operativo previo", [
                "Duplicidad de informacion",
                "Regularizacion tardia",
                "Poca trazabilidad documental",
                "Dificultad para seguimiento gerencial",
            ]),
            "key": "El proyecto ataca tanto el registro de lecturas como el control integral de la operacion."
        },
        {
            "title": "Arquitectura funcional de la solucion",
            "subtitle": "Vista end-to-end del flujo operativo",
            "flow": True,
            "bullets": [
                "Unifica acceso, padron, lectura automatica, carga manual y exportacion.",
                "Permite trazabilidad por periodo, cuadrilla, suministro, estado y origen.",
            ],
            "visual": ("Capturas para esta slide", [
                "Captura 02: Header y navegacion principal",
                "Captura 03: Selector de periodo y buscador",
            ]),
            "key": "La solucion organiza el proceso de extremo a extremo dentro de un solo flujo operativo."
        },
        {
            "title": "Panel de control y seguimiento gerencial",
            "subtitle": "Control en tiempo real del estado del periodo",
            "kpis": [
                ("TOTAL DEL PERIODO", "1109", COLOR_NAVY),
                ("PENDIENTES", "290", COLOR_ORANGE),
                ("LEIDOS SIN PARAMETROS", "0", COLOR_MUTED),
                ("REALIZADOS", "819", COLOR_GREEN),
            ],
            "bullets": [
                "Consulta por periodo y busqueda por suministro.",
                "Filtros por estado para priorizacion operativa.",
                "Listado navegable con acceso directo al detalle del suministro.",
                "Exportacion del padron operativo a Excel.",
            ],
            "visual": ("Evidencia visual sugerida", [
                "Captura 04: KPI del periodo",
                "Captura 05: Listado del padron",
            ]),
            "key": "Gerencia puede visualizar volumen, avance y excepciones del periodo desde una sola pantalla."
        },
        {
            "title": "Gestion del padron ENOSA",
            "subtitle": "Base maestra que ordena la operacion mensual",
            "bullets": [
                "Carga de padron por periodo con validacion de duplicados.",
                "Deteccion y control del periodo aplicable.",
                "Borrado controlado por mes con confirmacion operativa.",
                "Conservacion de estructura maestra para toda la trazabilidad posterior.",
            ],
            "visual": ("Evidencia visual sugerida", [
                "Captura 07: Modulo Cargar padron ENOSA",
            ]),
            "key": "El padron deja de ser un archivo aislado y se convierte en la base controlada del proceso."
        },
        {
            "title": "Procesamiento de lecturas por cuadrilla",
            "subtitle": "Automatizacion del flujo documental y operativo de campo",
            "bullets": [
                "Procesamiento individual por cuadrilla o ejecucion total.",
                "Control de procesados, omitidos, errores y archivos especiales.",
                "Organizacion de procesados por periodo y cuadrilla.",
                "Descarga de archivos sin parametros para su regularizacion posterior.",
            ],
            "visual": ("Evidencia visual sugerida", [
                "Captura 08: Modulo Procesar lecturas",
            ]),
            "key": "La operacion de campo pasa a ser medible, trazable y ordenada por equipo de trabajo."
        },
        {
            "title": "Soporte multi-formato y tratamiento de excepciones",
            "subtitle": "Adaptacion a la realidad tecnica de los equipos y archivos recibidos",
            "table": [
                ["Formato", "Tratamiento", "Resultado operativo", "Estado"],
                [".txt Metcom", "Lectura automatica de suministro y OBIS", "Parametros actuales y previos", "REALIZADO"],
                [".RG Alphaset", "Extraccion de 12 parametros", "Instantanea + previa", "REALIZADO"],
                [".msr / .RP3", "Lectura de suministro sin parametros", "Consolidacion para regularizacion", "LEIDO SIN PARAMETROS"],
            ],
            "bullets": [
                "Se normaliza el suministro incluso con variaciones como ceros iniciales.",
                "Se diferencia lectura completa versus lectura sin parametros para no perder trazabilidad.",
            ],
            "key": "El sistema absorbe formatos heterogeneos sin romper el control operativo."
        },
        {
            "title": "Detalle de suministro y trazabilidad puntual",
            "subtitle": "Consulta operativa de cada caso en el periodo",
            "bullets": [
                "Vista compacta de parametros, estado, origen, observacion y contexto del suministro.",
                "Separacion entre lectura previa e instantanea cuando aplica.",
                "Revision puntual para soporte operativo, validacion y auditoria.",
            ],
            "visual": ("Evidencia visual sugerida", [
                "Captura 06: Detalle del suministro",
            ]),
            "key": "Cada suministro puede revisarse de forma individual con contexto suficiente para tomar accion."
        },
        {
            "title": "Carga manual individual",
            "subtitle": "Ruta formal para regularizacion caso por caso",
            "bullets": [
                "Busqueda de pendientes por periodo y suministro.",
                "Registro individual de parametros y observacion.",
                "Cambio de estado y trazabilidad inmediata del caso.",
                "Soporte para correccion operativa sin salir de la plataforma.",
            ],
            "visual": ("Evidencia visual sugerida", [
                "Captura 09: Ingreso de carga manual",
                "Captura 12: Registrar carga",
            ]),
            "key": "Las excepciones no detienen el cierre: cuentan con un flujo de regularizacion controlado."
        },
        {
            "title": "Carga masiva por plantilla",
            "subtitle": "Escalabilidad para regularizacion de volumen",
            "bullets": [
                "Descarga de plantilla segun filtro y pendientes del periodo.",
                "Subida de Excel con validaciones por suministro y periodo.",
                "Registro masivo de parametros desde plataforma.",
                "Uso del valor visible del Excel para evitar distorsiones de precision.",
            ],
            "visual": ("Evidencia visual sugerida", [
                "Captura 10: Carga masiva por plantilla",
                "Captura 11: Pendientes del periodo",
            ]),
            "key": "La plataforma combina control individual y capacidad de carga masiva sin perder trazabilidad."
        },
        {
            "title": "Control de origen: TELEMETRIA vs CAMPO",
            "subtitle": "Clasificacion operativa y capacidad de correccion",
            "bullets": [
                "La carga masiva por plantilla clasifica automaticamente como TELEMETRIA.",
                "Los demas casos del periodo pueden quedar clasificados como CAMPO.",
                "El origen puede corregirse directamente desde la plataforma con edicion controlada.",
                "La misma vista permite ajustar parametros y observacion cuando se requiere.",
            ],
            "visual": ("Evidencia visual sugerida", [
                "Captura 06: Detalle del suministro con icono de edicion",
            ]),
            "key": "No solo se registra informacion: tambien se clasifica, corrige y formaliza el origen de cada lectura."
        },
        {
            "title": "Exportacion, consolidacion y evidencia documental",
            "subtitle": "Cierre operativo con salida util para gestion",
            "bullets": [
                "Exportacion del padron a Excel con criterios de negocio y ordenamiento de conceptos.",
                "Llenado de LecturaO con la carga disponible.",
                "ZIP por cuadrilla para archivos .msr y .RP3 con indice y plantilla de regularizacion.",
                "Conservacion estructurada de procesados por periodo y cuadrilla.",
            ],
            "visual": ("Salida operativa esperada", [
                "Excel del padron consolidado",
                "ZIP por cuadrilla",
                "Indice para regularizacion",
                "Trazabilidad documental en Drive",
            ]),
            "key": "La solucion no termina en la captura: genera salidas utiles para control, cierre y evidencia."
        },
        {
            "title": "Beneficios gerenciales y KPIs sugeridos",
            "subtitle": "Valor esperado para Gerencia General",
            "bullets": [
                "Mayor visibilidad del avance del periodo y del estado real de la operacion.",
                "Menor dependencia de cruces manuales dispersos.",
                "Mayor trazabilidad por suministro, archivo, cuadrilla y origen.",
                "Mejor capacidad de regularizacion y seguimiento de excepciones.",
            ],
            "visual": ("KPIs recomendados", [
                "% realizado por periodo",
                "% pendientes por periodo",
                "% TELEMETRIA vs CAMPO",
                "% leidos sin parametros",
                "Tiempo de cierre por cuadrilla",
                "Numero de correcciones manuales",
            ]),
            "key": "El sistema crea una base medible para gestionar productividad, excepciones y cierre operativo."
        },
        {
            "title": "Proximos pasos recomendados",
            "subtitle": "Ruta sugerida para consolidacion y escalamiento",
            "bullets": [
                "Formalizar rutina de seguimiento semanal por periodo y cuadrilla.",
                "Definir tablero gerencial de KPIs de cierre.",
                "Consolidar version estable para uso operativo continuo.",
                "Evaluar integracion futura con reportes historicos y control comparativo mensual.",
            ],
            "visual": ("Propuesta de cierre", [
                "Adopcion operativa",
                "Seguimiento gerencial",
                "Mejora continua",
                "Escalamiento funcional",
            ]),
            "key": "La siguiente etapa es institucionalizar el uso del sistema como herramienta de control operativo."
        },
    ]

    for index, data in enumerate(slides_data, start=1):
        slide = prs.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, data["title"], data.get("subtitle"))

        if index == 1:
            add_bullets(slide, Inches(0.8), Inches(1.75), Inches(6.0), Inches(2.1), data["bullets"], size=18)
            add_visual_placeholder(slide, Inches(7.55), Inches(1.55), Inches(4.9), Inches(3.0),
                                   data["visual"][0], data["visual"][1])
        elif data.get("flow"):
            add_process_flow(slide)
            add_bullets(slide, Inches(0.8), Inches(3.25), Inches(6.2), Inches(1.2), data["bullets"], size=16)
            add_visual_placeholder(slide, Inches(7.55), Inches(1.65), Inches(4.9), Inches(2.5),
                                   data["visual"][0], data["visual"][1])
        elif data.get("kpis"):
            kpi_x = [0.8, 3.15, 5.5, 7.85]
            for i, (title, value, color) in enumerate(data["kpis"]):
                add_kpi_card(slide, Inches(kpi_x[i]), Inches(1.65), Inches(2.15), Inches(1.2), title, value, color)
            add_bullets(slide, Inches(0.8), Inches(3.15), Inches(6.25), Inches(2.2), data["bullets"], size=16)
            add_visual_placeholder(slide, Inches(7.55), Inches(3.0), Inches(4.9), Inches(2.3),
                                   data["visual"][0], data["visual"][1])
        elif data.get("table"):
            add_table_matrix(slide, Inches(0.75), Inches(1.65), Inches(8.2), Inches(2.55), data["table"])
            add_bullets(slide, Inches(0.85), Inches(4.45), Inches(7.8), Inches(1.15), data["bullets"], size=15)
        else:
            add_bullets(slide, Inches(0.8), Inches(1.6), Inches(6.25), Inches(3.6), data["bullets"], size=17)
            add_visual_placeholder(slide, Inches(7.55), Inches(1.55), Inches(4.9), Inches(3.15),
                                   data["visual"][0], data["visual"][1])

        add_key_message(slide, data["key"])
        add_footer(slide, index)

    prs.save(OUTPUT_FILE)


if __name__ == "__main__":
    build_presentation()
    print(f"PPT generado: {OUTPUT_FILE.resolve()}")
